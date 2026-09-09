"""
Term sheet structured-data extraction.

Pipeline: .docx (text + tables) ->OPEN AI LLM (strict JSON) -> validated dict.

Usage:
    python termsheet_extractor.py test_termsheet.docx

Or as a library:
    from termsheet_extractor import extract_termsheet
    data = extract_termsheet("test_termsheet.docx")
"""

import json
import os
import sys

import docx
from groq import Groq

# The only fields we promise to return. Anything the document doesn't
# address comes back as None rather than a guess.
REQUIRED_FIELDS = [
    "partner_name",
    "contract_duration_months",
    "revenue_share",
    "revenue_share_details",
    "minimum_volume",
    "exclusivity",
    "payment_terms",
    "termination_terms",
    "renewal_terms",
    "IP_Ownership",
]

DEFAULT_MODEL = "openai/gpt-oss-120b"

EXTRACTION_PROMPT_TEMPLATE = """You are a contracts analyst extracting structured data from a partnership term sheet.

Read the DOCUMENT below and return a single JSON object with EXACTLY these keys:

- "partner_name": string or null — the name of the partner/counterparty company.
- "contract_duration_months": integer or null — the initial contract term, converted to months (e.g. "2 years" -> 24). If a duration is stated but not a fixed length (e.g. "until terminated"), use null.
- "revenue_share": string or null — the revenue share / commission arrangement, quoted or closely paraphrased from the document (e.g. "15% of net revenue").
 - "revenue_share_details": string or null -  the existing narrative text — how it's calculated, what it's a percentage of, any tiers/conditions
- "minimum_volume": string or null — any minimum volume, minimum order quantity, or minimum spend commitment.
- "exclusivity": string or null — the exclusivity arrangement. If the document explicitly states the relationship is non-exclusive, return "Non-exclusive" (not null). Use null only if exclusivity is not addressed at all.
- "payment_terms": string or null — payment schedule/timing terms (e.g. "Net 30 from invoice date").
- "termination_terms": string or null — conditions or notice period for termination.
- "renewal_terms": string or null — how/whether the agreement renews (e.g. "auto-renews annually unless 60 days' notice given").
- "IP_Ownership": string or null - How is the IP ownership split between the parties (e.g. "Company A owns their IP and any derivative work is owned by company B")

Rules:
- Use ONLY information present in the DOCUMENT. Do not infer or guess values that aren't stated.
- If a field is genuinely not addressed in the document, its value MUST be null — never invent a plausible-sounding value.
- Return ONLY the JSON object. No markdown fences, no commentary, no extra keys.

DOCUMENT:
\"\"\"
{document_text}
\"\"\"
"""

RETRY_SUFFIX = """

Your previous response could not be parsed as valid JSON. Return ONLY a single valid JSON object
with exactly the keys described above — no markdown code fences, no explanation, no trailing text.
"""


class ExtractionError(Exception):
    """Raised when the LLM response can't be turned into valid structured JSON."""


def extract_docx_content(docx_path) -> str:
    """Pull paragraph and table text out of a .docx file, in document order-ish
    (python-docx doesn't give true interleaved order without walking XML, but
    paragraphs-then-tables is sufficient for a term sheet's field content)."""
    document = docx.Document(docx_path)

    parts = []

    paragraph_text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    if paragraph_text:
        parts.append(paragraph_text)

    for i, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"[Table {i}]\n" + "\n".join(rows))

    return "\n\n".join(parts)


def _coerce_duration_months(value):
    """Best-effort coercion of contract_duration_months to int or None,
    without ever fabricating a value that wasn't in the model's output."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.isdigit():
            return int(stripped)
    return None


def _parse_json_response(raw_text: str):
    """Parse the model's response as JSON, tolerating an accidental markdown fence."""
    text = raw_text.strip()
    if text.startswith("```"):
        text = text.strip("`")
        if text.lstrip().lower().startswith("json"):
            text = text.lstrip()[4:]
        text = text.strip()
    return json.loads(text)


def _validate_and_fill(parsed: dict) -> dict:
    if not isinstance(parsed, dict):
        raise ExtractionError(f"Expected a JSON object, got {type(parsed).__name__}")

    result = {}
    for field in REQUIRED_FIELDS:
        result[field] = parsed.get(field, None)

    result["contract_duration_months"] = _coerce_duration_months(
        result["contract_duration_months"]
    )

    # Normalize empty strings / obvious "not specified" filler to null rather
    # than guessing, but never invent a value that wasn't returned.
    for field in REQUIRED_FIELDS:
        if isinstance(result[field], str):
            cleaned = result[field].strip()
            if not cleaned or cleaned.lower() in ("null", "none", "n/a", "not specified", "not mentioned"):
                result[field] = None
            else:
                result[field] = cleaned

    return result


def extract_termsheet(docx_path, groq_client: Groq = None, model: str = DEFAULT_MODEL) -> dict:
    """
    Extract structured term sheet fields from a .docx file.

    Returns a dict with exactly REQUIRED_FIELDS as keys. Any field not
    addressed in the document is None.

    Raises ExtractionError if the LLM never returns parseable JSON.
    """
    if groq_client is None:
        api_key = os.getenv("GROQ_API_KEY")
        if not api_key:
            raise ExtractionError("GROQ_API_KEY is not set and no groq_client was provided")
        groq_client = Groq(api_key=api_key)

    document_text = extract_docx_content(docx_path)
    if not document_text.strip():
        raise ExtractionError(f"No extractable text found in {docx_path}")

    prompt = EXTRACTION_PROMPT_TEMPLATE.format(document_text=document_text)

    last_raw = None
    for attempt in range(2):  # one retry if the first response isn't valid JSON
        response = groq_client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"},
            temperature=0,
        )
        last_raw = response.choices[0].message.content

        try:
            parsed = _parse_json_response(last_raw)
        except (json.JSONDecodeError, ValueError):
            prompt = EXTRACTION_PROMPT_TEMPLATE.format(document_text=document_text) + RETRY_SUFFIX
            continue

        return _validate_and_fill(parsed)

    raise ExtractionError(
        f"Model did not return valid JSON after retry. Last raw response:\n{last_raw}"
    )


if __name__ == "__main__":
    from dotenv import load_dotenv

    load_dotenv()

    if len(sys.argv) < 2:
        print("Usage: python termsheet_extractor.py <path-to-termsheet.docx>")
        sys.exit(1)

    path = sys.argv[1]
    try:
        data = extract_termsheet(path)
    except ExtractionError as e:
        print(f"Extraction failed: {e}")
        sys.exit(1)

    print(json.dumps(data, indent=2))
